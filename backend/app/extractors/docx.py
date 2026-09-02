import docx
import io
from typing import List
from fastapi import HTTPException, status

from app.extractors.base import BaseExtractor, ExtractedBlock, ExtractedDocument, SourceTypeEnum
from app.extractors.normalizer import normalize_text, build_full_text

class DOCXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse DOCX document: {str(e)}"
            )

        blocks: List[ExtractedBlock] = []
        seq_num = 1
        para_idx = 1

        # Extract paragraphs
        for p in doc.paragraphs:
            norm_text = normalize_text(p.text)
            if norm_text:
                style_name = p.style.name if p.style else None
                blocks.append(
                    ExtractedBlock(
                        sequence_number=seq_num,
                        source_type=SourceTypeEnum.PARAGRAPH,
                        source_index=para_idx,
                        text=norm_text,
                        metadata_json={"paragraph_index": para_idx, "style": style_name}
                    )
                )
                seq_num += 1
                para_idx += 1

        # Extract tables
        for table_idx, table in enumerate(doc.tables, start=1):
            table_rows = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                if any(row_cells):
                    table_rows.append(" | ".join(row_cells))
            if table_rows:
                table_text = normalize_text("\n".join(table_rows))
                if table_text:
                    blocks.append(
                        ExtractedBlock(
                            sequence_number=seq_num,
                            source_type=SourceTypeEnum.PARAGRAPH,
                            source_index=para_idx,
                            text=table_text,
                            metadata_json={"table_index": table_idx}
                        )
                    )
                    seq_num += 1
                    para_idx += 1

        if not blocks:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="DOCX file contains no extractable text content."
            )

        full_text = build_full_text(blocks)
        return ExtractedDocument(
            full_text=full_text,
            blocks=blocks,
            source_unit_count=para_idx - 1,
            character_count=len(full_text)
        )
