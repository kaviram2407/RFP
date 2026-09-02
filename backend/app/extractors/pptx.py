import pptx
import io
from typing import List
from fastapi import HTTPException, status

from app.extractors.base import BaseExtractor, ExtractedBlock, ExtractedDocument, SourceTypeEnum
from app.extractors.normalizer import normalize_text, build_full_text

class PPTXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        try:
            prs = pptx.Presentation(io.BytesIO(file_bytes))
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse PPTX presentation: {str(e)}"
            )

        blocks: List[ExtractedBlock] = []
        seq_num = 1
        slide_count = len(prs.slides)

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            slide_texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            slide_texts.append(text)

            if slide_texts:
                combined_text = normalize_text("\n".join(slide_texts))
                if combined_text:
                    blocks.append(
                        ExtractedBlock(
                            sequence_number=seq_num,
                            source_type=SourceTypeEnum.SLIDE,
                            source_index=slide_idx,
                            text=combined_text,
                            metadata_json={
                                "slide_number": slide_idx,
                                "slide_title": slide_title
                            }
                        )
                    )
                    seq_num += 1

        if not blocks:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="PPTX presentation contains no extractable text content."
            )

        full_text = build_full_text(blocks)
        return ExtractedDocument(
            full_text=full_text,
            blocks=blocks,
            source_unit_count=slide_count,
            character_count=len(full_text)
        )
