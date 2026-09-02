import openpyxl
import io
from typing import List
from fastapi import HTTPException, status

from app.extractors.base import BaseExtractor, ExtractedBlock, ExtractedDocument, SourceTypeEnum
from app.extractors.normalizer import normalize_text, build_full_text

class XLSXExtractor(BaseExtractor):
    def extract(self, file_bytes: bytes) -> ExtractedDocument:
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Failed to parse XLSX workbook: {str(e)}"
            )

        blocks: List[ExtractedBlock] = []
        seq_num = 1
        sheet_count = len(wb.sheetnames)

        for sheet_idx, sheet_name in enumerate(wb.sheetnames, start=1):
            sheet = wb[sheet_name]
            sheet_rows = []
            
            for row in sheet.iter_rows(values_only=True):
                # Filter out empty cell values
                cell_values = [str(val).strip() for val in row if val is not None and str(val).strip() != ""]
                if cell_values:
                    sheet_rows.append(" | ".join(cell_values))

            if sheet_rows:
                sheet_text = normalize_text("\n".join(sheet_rows))
                if sheet_text:
                    blocks.append(
                        ExtractedBlock(
                            sequence_number=seq_num,
                            source_type=SourceTypeEnum.SHEET,
                            source_index=sheet_idx,
                            text=sheet_text,
                            metadata_json={"sheet_name": sheet_name, "sheet_index": sheet_idx}
                        )
                    )
                    seq_num += 1

        wb.close()

        if not blocks:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail="XLSX workbook contains no extractable populated cell content."
            )

        full_text = build_full_text(blocks)
        return ExtractedDocument(
            full_text=full_text,
            blocks=blocks,
            source_unit_count=sheet_count,
            character_count=len(full_text)
        )
