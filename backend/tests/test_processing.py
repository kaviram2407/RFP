import pytest
from unittest.mock import patch, MagicMock
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import io
import fitz
import docx
import openpyxl
import pptx

from app.main import app
from app.db.base import Base
from app.api.deps import get_db
from app.models.organization import Organization
from app.models.user import User, RoleEnum
from app.models.rfp_project import RFPProject
from app.models.rfp_document import (
    RFPDocument,
    DocumentVersion,
    DocumentContent,
    DocumentContentBlock,
    DocumentTypeEnum,
    DocumentStatusEnum,
    ProcessingStatusEnum,
    SourceTypeEnum,
)
from app.extractors.pdf import PDFExtractor
from app.extractors.docx import DOCXExtractor
from app.extractors.xlsx import XLSXExtractor
from app.extractors.pptx import PPTXExtractor
from app.extractors.normalizer import normalize_text, build_full_text
from app.services.document_processing import process_document_version
from app.core.security import get_password_hash

SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

Base.metadata.create_all(bind=engine)

def override_get_db():
    try:
        db = TestingSessionLocal()
        yield db
    finally:
        db.close()

@pytest.fixture
def proc_db():
    app.dependency_overrides[get_db] = override_get_db
    db = TestingSessionLocal()
    for table in reversed(Base.metadata.sorted_tables):
        db.execute(table.delete())
    db.commit()

    orgA = Organization(name="Org A", slug="org-a")
    orgB = Organization(name="Org B", slug="org-b")
    db.add_all([orgA, orgB])
    db.commit()

    user_a = User(
        email="product_a@orga.com",
        full_name="Product A",
        password_hash=get_password_hash("password123"),
        organization_id=orgA.id,
        role=RoleEnum.PRODUCT_TEAM,
    )
    user_b = User(
        email="product_b@orgb.com",
        full_name="Product B",
        password_hash=get_password_hash("password123"),
        organization_id=orgB.id,
        role=RoleEnum.PRODUCT_TEAM,
    )
    db.add_all([user_a, user_b])
    db.commit()

    projA = RFPProject(name="Project A", reference_number="REF-A", organization_id=orgA.id, created_by_id=user_a.id)
    projB = RFPProject(name="Project B", reference_number="REF-B", organization_id=orgB.id, created_by_id=user_b.id)
    db.add_all([projA, projB])
    db.commit()

    docA = RFPDocument(name="Spec.pdf", document_type=DocumentTypeEnum.PDF, organization_id=orgA.id, rfp_project_id=projA.id, created_by_id=user_a.id)
    docB = RFPDocument(name="SpecB.pdf", document_type=DocumentTypeEnum.PDF, organization_id=orgB.id, rfp_project_id=projB.id, created_by_id=user_b.id)
    db.add_all([docA, docB])
    db.commit()

    verA = DocumentVersion(
        organization_id=orgA.id,
        rfp_document_id=docA.id,
        created_by_id=user_a.id,
        version_number=1,
        original_filename="Spec.pdf",
        storage_key=f"organizations/{orgA.id}/rfp-projects/{projA.id}/documents/{docA.id}/versions/v1",
        content_type="application/pdf",
        file_size_bytes=1024,
        checksum_sha256="dummyhash",
        processing_status=ProcessingStatusEnum.PENDING,
    )
    verB = DocumentVersion(
        organization_id=orgB.id,
        rfp_document_id=docB.id,
        created_by_id=user_b.id,
        version_number=1,
        original_filename="SpecB.pdf",
        storage_key=f"organizations/{orgB.id}/rfp-projects/{projB.id}/documents/{docB.id}/versions/v1",
        content_type="application/pdf",
        file_size_bytes=1024,
        checksum_sha256="dummyhashb",
        processing_status=ProcessingStatusEnum.PENDING,
    )
    db.add_all([verA, verB])
    db.commit()

    docA.current_version_id = verA.id
    docB.current_version_id = verB.id
    db.commit()

    yield {
        "db": db,
        "orgA": orgA,
        "orgB": orgB,
        "user_a": user_a,
        "user_b": user_b,
        "projA": projA,
        "docA": docA,
        "verA": verA,
        "projB": projB,
        "docB": docB,
        "verB": verB,
    }
    db.close()

def get_token(client: TestClient, email: str):
    res = client.post("/auth/login", data={"username": email, "password": "password123"})
    return res.json()["access_token"]

# --- Extractor Unit Tests ---

def test_pdf_extractor():
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Vendor must provide 24x7 support.")
    pdf_bytes = doc.tobytes()
    doc.close()

    extractor = PDFExtractor()
    extracted = extractor.extract(pdf_bytes)

    assert extracted.source_unit_count == 1
    assert "24x7 support" in extracted.full_text
    assert len(extracted.blocks) == 1
    assert extracted.blocks[0].source_type == SourceTypeEnum.PAGE
    assert extracted.blocks[0].source_index == 1

def test_docx_extractor():
    doc = docx.Document()
    doc.add_heading("Technical Requirements", level=1)
    doc.add_paragraph("All data must be encrypted at rest using AES-256.")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    extractor = DOCXExtractor()
    extracted = extractor.extract(docx_bytes)

    assert extracted.source_unit_count >= 1
    assert "AES-256" in extracted.full_text
    assert extracted.blocks[0].source_type == SourceTypeEnum.PARAGRAPH

def test_xlsx_extractor():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Compliance Matrix"
    ws.append(["Req ID", "Description", "Compliant"])
    ws.append(["REQ-01", "SOC2 Type II Certification", "Yes"])

    stream = io.BytesIO()
    wb.save(stream)
    xlsx_bytes = stream.getvalue()

    extractor = XLSXExtractor()
    extracted = extractor.extract(xlsx_bytes)

    assert extracted.source_unit_count == 1
    assert "SOC2 Type II" in extracted.full_text
    assert extracted.blocks[0].source_type == SourceTypeEnum.SHEET
    assert extracted.blocks[0].metadata_json["sheet_name"] == "Compliance Matrix"

def test_pptx_extractor():
    prs = pptx.Presentation()
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    txBox = slide.shapes.add_textbox(0, 0, 100, 100)
    tf = txBox.text_frame
    tf.text = "Slide 1 Architecture Overview"

    stream = io.BytesIO()
    prs.save(stream)
    pptx_bytes = stream.getvalue()

    extractor = PPTXExtractor()
    extracted = extractor.extract(pptx_bytes)

    assert extracted.source_unit_count == 1
    assert "Architecture Overview" in extracted.full_text
    assert extracted.blocks[0].source_type == SourceTypeEnum.SLIDE

# --- Normalizer Tests ---

def test_normalizer_cleanup():
    raw = "Header \r\n\n\n Paragraph line 1   \n   Paragraph line 2  "
    norm = normalize_text(raw)
    assert "\r" not in norm
    assert "Paragraph line 1\nParagraph line 2" in norm

# --- Processing Lifecycle & Service Tests ---

@patch("app.services.document_processing.storage_service")
def test_process_document_version_success(mock_storage, proc_db):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "Sample PDF Text Content")
    pdf_bytes = doc.tobytes()
    doc.close()

    mock_storage.bucket_name = "rfp-documents"
    mock_storage.s3_client.get_object.return_value = {"Body": io.BytesIO(pdf_bytes)}

    db = proc_db["db"]
    ver_id = proc_db["verA"].id

    res_ver = process_document_version(db, ver_id)
    assert res_ver.processing_status == ProcessingStatusEnum.COMPLETED
    assert res_ver.processing_completed_at is not None
    assert res_ver.processing_error is None

    # Check database persistence
    content = db.query(DocumentContent).filter(DocumentContent.document_version_id == ver_id).first()
    assert content is not None
    assert "Sample PDF Text Content" in content.full_text

    blocks = db.query(DocumentContentBlock).filter(DocumentContentBlock.document_version_id == ver_id).all()
    assert len(blocks) >= 1
    assert blocks[0].source_type == SourceTypeEnum.PAGE

@patch("app.services.document_processing.storage_service")
def test_process_document_version_idempotency_retry(mock_storage, proc_db):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "First Run Text")
    pdf_bytes = doc.tobytes()
    doc.close()

    mock_storage.bucket_name = "rfp-documents"
    mock_storage.s3_client.get_object.return_value = {"Body": io.BytesIO(pdf_bytes)}

    db = proc_db["db"]
    ver_id = proc_db["verA"].id

    # Run 1
    process_document_version(db, ver_id)
    blocks1_count = db.query(DocumentContentBlock).filter(DocumentContentBlock.document_version_id == ver_id).count()

    # Run 2 (Retry)
    process_document_version(db, ver_id)
    blocks2_count = db.query(DocumentContentBlock).filter(DocumentContentBlock.document_version_id == ver_id).count()

    # Should replace, not duplicate
    assert blocks1_count == blocks2_count
    assert db.query(DocumentContent).filter(DocumentContent.document_version_id == ver_id).count() == 1

@patch("app.services.document_processing.storage_service")
def test_process_document_version_failure_handling(mock_storage, proc_db):
    mock_storage.bucket_name = "rfp-documents"
    mock_storage.s3_client.get_object.side_effect = Exception("R2 Connection Timeout")

    db = proc_db["db"]
    ver_id = proc_db["verA"].id

    res_ver = process_document_version(db, ver_id)
    assert res_ver.processing_status == ProcessingStatusEnum.FAILED
    assert "R2 Connection Timeout" in res_ver.processing_error

# --- API Endpoints & Tenant Isolation Tests ---

@patch("app.services.document_processing.storage_service")
def test_processing_api_and_tenant_isolation(mock_storage, proc_db):
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 50), "API Test Content")
    pdf_bytes = doc.tobytes()
    doc.close()

    mock_storage.bucket_name = "rfp-documents"
    mock_storage.s3_client.get_object.return_value = {"Body": io.BytesIO(pdf_bytes)}

    client = TestClient(app)
    token_a = get_token(client, "product_a@orga.com")
    token_b = get_token(client, "product_b@orgb.com")

    proj_id_a = proc_db["projA"].id
    doc_id_a = proc_db["docA"].id
    ver_id_a = proc_db["verA"].id

    # Trigger processing via API
    res_trigger = client.post(
        f"/api/v1/rfp-projects/{proj_id_a}/documents/{doc_id_a}/versions/{ver_id_a}/process",
        headers={"Authorization": f"Bearer {token_a}"}
    )
    assert res_trigger.status_code == 200
    assert res_trigger.json()["status"] == "COMPLETED"

    # Get processing status
    res_status = client.get(
        f"/api/v1/rfp-projects/{proj_id_a}/documents/{doc_id_a}/versions/{ver_id_a}/processing",
        headers={"Authorization": f"Bearer {token_a}"}
    )
    assert res_status.status_code == 200
    assert res_status.json()["status"] == "COMPLETED"

    # Get extracted content
    res_content = client.get(
        f"/api/v1/rfp-projects/{proj_id_a}/documents/{doc_id_a}/versions/{ver_id_a}/content",
        headers={"Authorization": f"Bearer {token_a}"}
    )
    assert res_content.status_code == 200
    assert "API Test Content" in res_content.json()["full_text"]

    # Cross-tenant access attempt by User B -> 404 Not Found
    res_cross_status = client.get(
        f"/api/v1/rfp-projects/{proj_id_a}/documents/{doc_id_a}/versions/{ver_id_a}/processing",
        headers={"Authorization": f"Bearer {token_b}"}
    )
    assert res_cross_status.status_code == 404
